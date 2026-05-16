# sql_chart_toolkit_v2.py
# -*- coding: utf-8 -*-

"""
SqlChartAthenaToolKitV2
=======================

Toolkit para:
- Executar SQL no Athena
- Gerar gráfico PNG
- Gerar Excel com os dados usados no gráfico
- Subir ambos no S3
- Retornar links autenticados (presigned URLs)

Suporta:
- pie
- bar
- column
- line
- area
- scatter
- indicator
- candlestick
- combo (bar/column + line)

Variáveis de ambiente esperadas:
- GRAPHS_S3_BASE
- EXCEL_S3_BASE
- ATHENA_OUTPUT_LOCATION
- AWS_ACCESS_KEY_ID
- AWS_SECRET_ACCESS_KEY
- AWS_DEFAULT_REGION
"""

import os
import re
import time
import json
import logging
from typing import Any, List, Dict, Optional, Callable, Tuple, Union

import boto3
from botocore.exceptions import ClientError

try:
    import pandas as pd
    import numpy as np
except Exception as e:
    raise ImportError(
        "Dependências ausentes. Instale: pandas, numpy. "
        f"Erro original: {e}"
    )

try:
    from openpyxl import load_workbook
    from openpyxl.styles import Font
except Exception as e:
    raise ImportError(
        "Dependências ausentes. Instale: openpyxl. "
        f"Erro original: {e}"
    )

logger = logging.getLogger(__name__)


class SqlChartAthenaToolKitV2:
    """
    Tool para gerar gráfico + Excel a partir de SQL Athena.
    """

    SUPPORTED_CHART_TYPES = {
        "pie",
        "bar",
        "column",
        "line",
        "area",
        "scatter",
        "indicator",
        "candlestick",
        "combo",
    }

    DEFAULT_COLORS = [
        "#1f77b4", "#ff7f0e", "#2ca02c", "#d62728", "#9467bd",
        "#8c564b", "#e377c2", "#7f7f7f", "#bcbd22", "#17becf"
    ]

    def __init__(
        self,
        athena_database: Optional[str] = None,
        graphs_s3_base: Optional[str] = None,
        excel_s3_base: Optional[str] = None,
        athena_output_location: Optional[str] = None,
        region_name: Optional[str] = None,
        aws_access_key_id: Optional[str] = None,
        aws_secret_access_key: Optional[str] = None,
        aws_session_token: Optional[str] = None,
        poll_interval: int = 2,
        max_rows: int = 100000,
        presign_expires_seconds: int = 604800,
    ) -> None:
        self.region_name = region_name or "us-east-1"
        self.aws_access_key_id = aws_access_key_id or ""
        self.aws_secret_access_key = aws_secret_access_key or ""
        self.aws_session_token = aws_session_token or ""

        self.athena_database = athena_database or ""
        self.athena_output_location = athena_output_location or ""
        self.graphs_s3_base = graphs_s3_base or ""
        self.excel_s3_base = excel_s3_base or ""

        self.poll_interval = int(poll_interval)
        self.max_rows = int(max_rows)
        self.presign_expires_seconds = int(presign_expires_seconds)

        if not self.aws_access_key_id or not self.aws_secret_access_key:
            raise ValueError("Credenciais AWS ausentes (AWS_ACCESS_KEY_ID/AWS_SECRET_ACCESS_KEY).")
        if not self.graphs_s3_base:
            raise ValueError("GRAPHS_S3_BASE ausente.")
        if not self.graphs_s3_base.startswith("s3://"):
            raise ValueError("GRAPHS_S3_BASE deve ser s3://bucket/prefix/")
        if not self.excel_s3_base:
            raise ValueError("EXCEL_S3_BASE ausente.")
        if not self.excel_s3_base.startswith("s3://"):
            raise ValueError("EXCEL_S3_BASE deve ser s3://bucket/prefix/")
        if not self.athena_output_location:
            raise ValueError("ATHENA_OUTPUT_LOCATION não configurado.")

        self.athena_client = boto3.client(
            "athena",
            region_name=self.region_name,
            aws_access_key_id=self.aws_access_key_id,
            aws_secret_access_key=self.aws_secret_access_key,
            aws_session_token=self.aws_session_token or None,
        )
        self.s3_client = boto3.client(
            "s3",
            region_name=self.region_name,
            aws_access_key_id=self.aws_access_key_id,
            aws_secret_access_key=self.aws_secret_access_key,
            aws_session_token=self.aws_session_token or None,
        )

        self._ensure_bucket_if_needed(self.athena_output_location)

        logger.info(
            "[SqlChartAthenaToolKitV2] Inicializado | Região=%s | AthenaOutput=%s | GraphBase=%s | ExcelBase=%s",
            self.region_name, self.athena_output_location, self.graphs_s3_base, self.excel_s3_base
        )

    # =========================================================================
    # Helpers básicos
    # =========================================================================

    @staticmethod
    def _extract_bucket_name(s3_uri: str) -> str:
        if not s3_uri.startswith("s3://"):
            raise ValueError(f"S3 URI inválida: {s3_uri}")
        return s3_uri.replace("s3://", "", 1).split("/", 1)[0]

    @staticmethod
    def _split_s3_uri(s3_uri: str) -> Tuple[str, str]:
        s3_uri = s3_uri.strip()
        if not s3_uri.startswith("s3://"):
            raise ValueError(f"URI S3 inválida: {s3_uri}")
        no_scheme = s3_uri.replace("s3://", "", 1)
        parts = no_scheme.split("/", 1)
        bucket = parts[0]
        prefix = parts[1] if len(parts) > 1 else ""
        return bucket, prefix

    @staticmethod
    def _slugify(texto: str) -> str:
        texto = (texto or "").strip().lower()
        texto = re.sub(r"[^\w\s-]", "", texto, flags=re.UNICODE)
        texto = re.sub(r"[\s-]+", "_", texto, flags=re.UNICODE)
        texto = re.sub(r"_+", "_", texto).strip("_")
        return texto or "arquivo"

    @staticmethod
    def _sheet_name(label: str) -> str:
        name = (label or "Planilha").strip()
        name = re.sub(r'[:\\/*?\[\]]', "_", name)
        name = name[:31].strip()
        return name or "Planilha"

    @staticmethod
    def _safe_float(x: Any) -> Optional[float]:
        if x is None:
            return None
        s = str(x).strip()
        if s == "" or s.lower() in ("nan", "none", "null"):
            return None

        if "," in s and "." in s:
            if s.rfind(",") > s.rfind("."):
                s = s.replace(".", "").replace(",", ".")
            else:
                s = s.replace(",", "")
        else:
            if "," in s and "." not in s:
                s = s.replace(",", ".")

        try:
            return float(s)
        except Exception:
            return None

    @staticmethod
    def _fmt_ptbr(v: Any, casas: int = 2) -> str:
        try:
            if v is None:
                return ""
            val = float(v)
            if np.isnan(val) or np.isinf(val):
                return ""
            s = f"{val:,.{casas}f}"
            return s.replace(",", "X").replace(".", ",").replace("X", ".")
        except Exception:
            return str(v)

    @staticmethod
    def _to_bool(v: Any, default: bool = False) -> bool:
        if v is None:
            return default
        if isinstance(v, bool):
            return v
        s = str(v).strip().lower()
        return s in ("1", "true", "t", "yes", "y", "sim")

    @staticmethod
    def _normalize_chart_type(chart_type: str) -> str:
        c = str(chart_type or "").strip().lower()
        aliases = {
            "pizza": "pie",
            "piechart": "pie",
            "barras": "bar",
            "barra": "bar",
            "colunas": "column",
            "coluna": "column",
            "linha": "line",
            "dispersao": "scatter",
            "indicador": "indicator",
            "candles": "candlestick",
            "candle": "candlestick",
            "combo_chart": "combo",
        }
        return aliases.get(c, c)

    def _ensure_bucket_if_needed(self, s3_uri: str) -> None:
        bucket_name = self._extract_bucket_name(s3_uri)
        try:
            self.s3_client.head_bucket(Bucket=bucket_name)
        except ClientError as e:
            code = e.response["Error"].get("Code", "")
            if code in ("404", "NoSuchBucket", "NotFound"):
                create_kwargs: Dict[str, Any] = {"Bucket": bucket_name}
                if self.region_name != "us-east-1":
                    create_kwargs["CreateBucketConfiguration"] = {"LocationConstraint": self.region_name}
                self.s3_client.create_bucket(**create_kwargs)
            else:
                raise

    # =========================================================================
    # Athena
    # =========================================================================

    def _execute_athena_query(self, query: str) -> pd.DataFrame:
        if not isinstance(query, str) or not query.strip():
            raise ValueError("Nenhuma consulta SQL válida foi fornecida.")
        if not self.athena_database:
            raise ValueError("athena_database não configurado.")
        if not self.athena_output_location:
            raise ValueError("ATHENA_OUTPUT_LOCATION não configurado.")

        response = self.athena_client.start_query_execution(
            QueryString=query,
            QueryExecutionContext={"Database": self.athena_database},
            ResultConfiguration={"OutputLocation": self.athena_output_location},
        )
        qid = response["QueryExecutionId"]

        while True:
            status = self.athena_client.get_query_execution(QueryExecutionId=qid)
            state = status["QueryExecution"]["Status"]["State"]

            if state in ("SUCCEEDED", "FAILED", "CANCELLED"):
                if state != "SUCCEEDED":
                    reason = status["QueryExecution"]["Status"].get("StateChangeReason", "Sem motivo detalhado")
                    raise RuntimeError(f"Query falhou. Estado: {state}. Motivo: {reason}")
                break

            time.sleep(self.poll_interval)

        paginator = self.athena_client.get_paginator("get_query_results")
        pages = paginator.paginate(QueryExecutionId=qid)

        rows: List[Dict[str, Any]] = []
        column_info = None

        for page in pages:
            if column_info is None:
                column_info = page["ResultSet"]["ResultSetMetadata"]["ColumnInfo"]

            for row in page["ResultSet"]["Rows"]:
                if (
                    row.get("Data")
                    and "VarCharValue" in row["Data"][0]
                    and row["Data"][0]["VarCharValue"] == column_info[0]["Name"]
                ):
                    continue

                row_dict: Dict[str, Any] = {}
                for col, cell in zip(column_info, row.get("Data", [])):
                    row_dict[col["Name"]] = cell.get("VarCharValue", None)

                rows.append(row_dict)

                if self.max_rows and len(rows) >= self.max_rows:
                    raise RuntimeError(
                        f"Resultado excedeu max_rows={self.max_rows}. Restrinja/agregue a SQL antes de gerar o gráfico."
                    )

        return pd.DataFrame(rows)

    # =========================================================================
    # Normalização
    # =========================================================================

    def _coerce_numeric_columns(self, df: pd.DataFrame, cols: List[str]) -> pd.DataFrame:
        out = df.copy()
        for c in cols:
            if c and c in out.columns:
                out[c] = out[c].apply(self._safe_float)
        return out

    @staticmethod
    def _coerce_datetime_column(df: pd.DataFrame, col: Optional[str]) -> pd.DataFrame:
        out = df.copy()
        if col and col in out.columns:
            out[col] = pd.to_datetime(out[col], errors="coerce")
        return out

    @staticmethod
    def _validate_columns(df: pd.DataFrame, required: List[str]) -> None:
        missing = [c for c in required if c and c not in df.columns]
        if missing:
            raise ValueError(f"Colunas obrigatórias ausentes no resultado: {missing}. Disponíveis: {list(df.columns)}")

    def _apply_top_n(
        self,
        df: pd.DataFrame,
        chart_type: str,
        top_n: Optional[int],
        sort_by: Optional[str],
        ascending: bool = False,
        value_col: Optional[str] = None,
        y_col: Optional[str] = None,
    ) -> pd.DataFrame:
        if not top_n or top_n <= 0:
            return df

        out = df.copy()

        chosen_sort = sort_by
        if not chosen_sort:
            if chart_type == "pie":
                chosen_sort = value_col
            else:
                chosen_sort = y_col

        if chosen_sort and chosen_sort in out.columns:
            out = out.sort_values(by=chosen_sort, ascending=ascending).head(int(top_n))

        return out

    def _pick_colors(self, payload: Dict[str, Any], n: int) -> List[str]:
        colors = payload.get("colors")
        if isinstance(colors, list) and len(colors) > 0:
            vals = [str(c).strip() for c in colors if str(c).strip()]
            if vals:
                if len(vals) >= n:
                    return vals[:n]
                rep = []
                while len(rep) < n:
                    rep.extend(vals)
                return rep[:n]
        rep = []
        while len(rep) < n:
            rep.extend(self.DEFAULT_COLORS)
        return rep[:n]

    # =========================================================================
    # Tendência
    # =========================================================================

    @staticmethod
    def _build_linear_trend(x_vals: np.ndarray, y_vals: np.ndarray) -> Optional[np.ndarray]:
        try:
            mask = np.isfinite(x_vals) & np.isfinite(y_vals)
            x = x_vals[mask]
            y = y_vals[mask]
            if len(x) < 2:
                return None
            coef = np.polyfit(x, y, 1)
            trend = np.poly1d(coef)(x_vals)
            return trend
        except Exception:
            return None

    # =========================================================================
    # Dados efetivamente usados no gráfico
    # =========================================================================

    def _prepare_chart_dataframe(self, df: pd.DataFrame, payload: Dict[str, Any]) -> pd.DataFrame:
        chart_type = self._normalize_chart_type(payload.get("chart_type"))
        x_col = payload.get("x_col")
        y_col = payload.get("y_col")
        y2_col = payload.get("y2_col")
        series_col = payload.get("series_col")
        category_col = payload.get("category_col")
        value_col = payload.get("value_col")
        size_col = payload.get("size_col")
        color_col = payload.get("color_col")
        open_col = payload.get("open_col")
        high_col = payload.get("high_col")
        low_col = payload.get("low_col")
        close_col = payload.get("close_col")
        top_n = payload.get("top_n")
        top_n = int(top_n) if top_n not in (None, "") else None
        sort_by = str(payload.get("sort_by") or "").strip() or None
        sort_ascending = self._to_bool(payload.get("sort_ascending"), default=False)

        work = df.copy()

        if chart_type == "pie":
            self._validate_columns(work, [category_col, value_col])
            work = work[[category_col, value_col]].copy()
            work = self._coerce_numeric_columns(work, [value_col])
            work = work.dropna(subset=[category_col, value_col])
            work[category_col] = work[category_col].astype(str)
            work = self._apply_top_n(work, chart_type, top_n, sort_by, sort_ascending, value_col=value_col)
            return work.reset_index(drop=True)

        if chart_type in ("bar", "column", "line", "area", "scatter"):
            req = [x_col, y_col]
            if series_col:
                req.append(series_col)
            if size_col:
                req.append(size_col)
            if color_col and color_col in work.columns:
                req.append(color_col)

            self._validate_columns(work, req)

            numeric_cols = [y_col]
            if chart_type == "scatter":
                numeric_cols = [x_col, y_col]
                if size_col:
                    numeric_cols.append(size_col)

            work = self._coerce_numeric_columns(work, [c for c in numeric_cols if c])
            if chart_type in ("line", "area"):
                work = self._coerce_datetime_column(work, x_col)

            keep_cols = [c for c in [x_col, y_col, series_col, size_col, color_col] if c and c in work.columns]
            work = work[keep_cols].copy()
            work = work.dropna(subset=[c for c in [x_col, y_col] if c in work.columns])

            if chart_type in ("bar", "column"):
                work = self._apply_top_n(work, chart_type, top_n, sort_by, sort_ascending, y_col=y_col)

            return work.reset_index(drop=True)

        if chart_type == "indicator":
            value_column = value_col or y_col
            if not value_column:
                raise ValueError("Para indicator, informe value_col ou y_col.")
            self._validate_columns(work, [value_column])

            keep_cols = [value_column]
            secondary_value_col = payload.get("secondary_value_col")
            if secondary_value_col and secondary_value_col in work.columns:
                keep_cols.append(secondary_value_col)

            work = work[keep_cols].copy()
            work = self._coerce_numeric_columns(work, keep_cols)
            work = work.dropna(subset=[value_column])
            return work.reset_index(drop=True)

        if chart_type == "candlestick":
            self._validate_columns(work, [x_col, open_col, high_col, low_col, close_col])
            keep_cols = [x_col, open_col, high_col, low_col, close_col]
            work = work[keep_cols].copy()
            work = self._coerce_datetime_column(work, x_col)
            work = self._coerce_numeric_columns(work, [open_col, high_col, low_col, close_col])
            work = work.dropna(subset=keep_cols).sort_values(by=x_col)
            return work.reset_index(drop=True)

        if chart_type == "combo":
            self._validate_columns(work, [x_col, y_col, y2_col])
            if series_col:
                self._validate_columns(work, [series_col])

            keep_cols = [c for c in [x_col, y_col, y2_col, series_col] if c and c in work.columns]
            work = work[keep_cols].copy()
            work = self._coerce_numeric_columns(work, [y_col, y2_col])
            work = self._coerce_datetime_column(work, x_col)
            work = work.dropna(subset=[x_col, y_col, y2_col])
            return work.reset_index(drop=True)

        raise ValueError(f"Tipo de gráfico não suportado: {chart_type}")

    # =========================================================================
    # Gráfico
    # =========================================================================

    def _create_chart_image(self, chart_df: pd.DataFrame, payload: Dict[str, Any]) -> str:
        import matplotlib.pyplot as plt
        import matplotlib.dates as mdates
        import matplotlib.ticker as mticker
        from matplotlib.patches import Rectangle

        chart_type = self._normalize_chart_type(payload.get("chart_type"))
        title = str(payload.get("title") or payload.get("label") or "Gráfico").strip()
        subtitle = str(payload.get("subtitle") or "").strip()

        x_col = payload.get("x_col")
        y_col = payload.get("y_col")
        y2_col = payload.get("y2_col")
        series_col = payload.get("series_col")
        category_col = payload.get("category_col")
        value_col = payload.get("value_col")
        size_col = payload.get("size_col")
        color_col = payload.get("color_col")
        open_col = payload.get("open_col")
        high_col = payload.get("high_col")
        low_col = payload.get("low_col")
        close_col = payload.get("close_col")

        horizontal = self._to_bool(payload.get("horizontal"), default=False)
        stacked = self._to_bool(payload.get("stacked"), default=False)
        trendline = self._to_bool(payload.get("trendline"), default=False)
        show_legend = self._to_bool(payload.get("show_legend"), default=True)
        show_labels = self._to_bool(payload.get("show_labels"), default=False)
        donut = self._to_bool(payload.get("donut"), default=False)
        figsize = payload.get("figsize") or [14, 8]
        dpi = int(payload.get("dpi") or 170)
        x_label = str(payload.get("x_label") or x_col or "").strip()
        y_label = str(payload.get("y_label") or y_col or "").strip()
        y2_label = str(payload.get("y2_label") or y2_col or "").strip()
        value_format = str(payload.get("value_format") or "number").strip().lower()
        secondary_value_format = str(payload.get("secondary_value_format") or value_format).strip().lower()

        try:
            width = float(figsize[0]) if isinstance(figsize, (list, tuple)) and len(figsize) >= 2 else 14.0
            height = float(figsize[1]) if isinstance(figsize, (list, tuple)) and len(figsize) >= 2 else 8.0
        except Exception:
            width, height = 14.0, 8.0

        def fmt_axis_value_factory(fmt_kind: str):
            def _fmt(v: Any, pos: Any = None) -> str:
                try:
                    val = float(v)
                except Exception:
                    return str(v)

                if fmt_kind == "percent":
                    return f"{val:,.2f}%".replace(",", "X").replace(".", ",").replace("X", ".")
                if fmt_kind == "currency":
                    return self._fmt_ptbr(val)
                if fmt_kind == "integer":
                    return self._fmt_ptbr(val, casas=0)
                if fmt_kind == "million":
                    return f"{self._fmt_ptbr(val / 1_000_000)} MM"
                return self._fmt_ptbr(val)
            return _fmt

        fig, ax = plt.subplots(figsize=(width, height))
        fig.patch.set_facecolor("white")
        ax.set_facecolor("white")

        # ---------------------------------------------------------------------
        # PIE
        # ---------------------------------------------------------------------
        if chart_type == "pie":
            colors = self._pick_colors(payload, len(chart_df))
            wedges, texts, autotexts = ax.pie(
                chart_df[value_col].values,
                labels=chart_df[category_col].values if show_labels else None,
                autopct="%1.1f%%",
                startangle=90,
                colors=colors,
            )
            ax.axis("equal")
            if donut:
                centre_circle = plt.Circle((0, 0), 0.60, fc="white")
                fig.gca().add_artist(centre_circle)

            if show_legend:
                ax.legend(
                    wedges,
                    chart_df[category_col].astype(str).tolist(),
                    title=category_col,
                    loc="center left",
                    bbox_to_anchor=(1, 0.5)
                )

        # ---------------------------------------------------------------------
        # BAR / COLUMN
        # ---------------------------------------------------------------------
        elif chart_type in ("bar", "column"):
            if series_col:
                pivot = chart_df.pivot_table(
                    index=x_col,
                    columns=series_col,
                    values=y_col,
                    aggfunc="sum",
                    fill_value=0
                ).reset_index()

                categories = pivot[x_col].astype(str).tolist()
                series_names = [c for c in pivot.columns if c != x_col]
                colors = self._pick_colors(payload, len(series_names))
                x = np.arange(len(categories))
                bottoms = np.zeros(len(categories), dtype=float)
                bar_width = 0.8 / max(len(series_names), 1)

                for i, sname in enumerate(series_names):
                    vals = pivot[sname].astype(float).values
                    if chart_type == "bar" and horizontal:
                        if stacked:
                            ax.barh(categories, vals, left=bottoms, label=str(sname), color=colors[i])
                            bottoms += vals
                        else:
                            ax.barh(x + (i * bar_width), vals, height=bar_width, label=str(sname), color=colors[i])
                    else:
                        if stacked:
                            ax.bar(categories, vals, bottom=bottoms, label=str(sname), color=colors[i])
                            bottoms += vals
                        else:
                            ax.bar(x + (i * bar_width), vals, width=bar_width, label=str(sname), color=colors[i])

                if not stacked and not (chart_type == "bar" and horizontal):
                    ax.set_xticks(x + (bar_width * max(len(series_names) - 1, 0) / 2))
                    ax.set_xticklabels(categories, rotation=45, ha="right")
                elif chart_type != "bar" or not horizontal:
                    ax.tick_params(axis="x", rotation=45)
            else:
                colors = self._pick_colors(payload, len(chart_df))
                if chart_type == "bar" and horizontal:
                    ax.barh(chart_df[x_col].astype(str), chart_df[y_col].astype(float), color=colors)
                else:
                    ax.bar(chart_df[x_col].astype(str), chart_df[y_col].astype(float), color=colors)
                    ax.tick_params(axis="x", rotation=45)

                if trendline and len(chart_df) >= 2:
                    xv = np.arange(len(chart_df), dtype=float)
                    yv = chart_df[y_col].astype(float).values
                    trend = self._build_linear_trend(xv, yv)
                    if trend is not None:
                        if chart_type == "bar" and horizontal:
                            ax.plot(trend, chart_df[x_col].astype(str), linestyle="--", linewidth=2, label="Tendência")
                        else:
                            ax.plot(chart_df[x_col].astype(str), trend, linestyle="--", linewidth=2, label="Tendência")

            if chart_type == "bar" and horizontal:
                ax.set_xlabel(y_label or "")
                ax.set_ylabel(x_label or "")
            else:
                ax.set_xlabel(x_label or "")
                ax.set_ylabel(y_label or "")

            ax.yaxis.set_major_formatter(mticker.FuncFormatter(fmt_axis_value_factory(value_format)))

        # ---------------------------------------------------------------------
        # LINE / AREA
        # ---------------------------------------------------------------------
        elif chart_type in ("line", "area"):
            if series_col:
                series_names = list(chart_df[series_col].dropna().astype(str).unique())
                colors = self._pick_colors(payload, len(series_names))

                for i, sname in enumerate(series_names):
                    sub = chart_df[chart_df[series_col].astype(str) == sname].copy().sort_values(by=x_col)
                    if chart_type == "line":
                        ax.plot(sub[x_col], sub[y_col].astype(float), marker="o", label=sname, color=colors[i])
                    else:
                        ax.plot(sub[x_col], sub[y_col].astype(float), marker="o", label=sname, color=colors[i])
                        ax.fill_between(sub[x_col], sub[y_col].astype(float), alpha=0.25, color=colors[i])
            else:
                work = chart_df.sort_values(by=x_col)
                color = self._pick_colors(payload, 1)[0]
                ax.plot(work[x_col], work[y_col].astype(float), marker="o", color=color, label=y_col if show_legend else None)
                if chart_type == "area":
                    ax.fill_between(work[x_col], work[y_col].astype(float), alpha=0.25, color=color)

                if trendline and len(work) >= 2:
                    x_num = np.arange(len(work), dtype=float)
                    y_num = work[y_col].astype(float).values
                    trend = self._build_linear_trend(x_num, y_num)
                    if trend is not None:
                        ax.plot(work[x_col], trend, linestyle="--", linewidth=2, label="Tendência")

            ax.set_xlabel(x_label or "")
            ax.set_ylabel(y_label or "")
            ax.yaxis.set_major_formatter(mticker.FuncFormatter(fmt_axis_value_factory(value_format)))

            if np.issubdtype(chart_df[x_col].dtype, np.datetime64):
                ax.xaxis.set_major_formatter(mdates.DateFormatter("%Y-%m-%d"))
                fig.autofmt_xdate()

        # ---------------------------------------------------------------------
        # SCATTER
        # ---------------------------------------------------------------------
        elif chart_type == "scatter":
            sizes = None
            if size_col and size_col in chart_df.columns:
                raw_sizes = chart_df[size_col].astype(float).values
                sizes = np.where(np.isfinite(raw_sizes), raw_sizes, 20.0)
                sizes = np.clip(sizes, 20, np.nanmax(sizes) if len(sizes) else 20)
                sizes = (sizes / np.nanmax(sizes)) * 400 if np.nanmax(sizes) > 0 else 60

            if color_col and color_col in chart_df.columns:
                if pd.api.types.is_numeric_dtype(chart_df[color_col]):
                    cvals = pd.to_numeric(chart_df[color_col], errors="coerce")
                    sc = ax.scatter(
                        chart_df[x_col].astype(float),
                        chart_df[y_col].astype(float),
                        s=sizes,
                        c=cvals,
                        alpha=0.75
                    )
                    fig.colorbar(sc, ax=ax, label=color_col)
                else:
                    groups = list(chart_df[color_col].astype(str).unique())
                    colors = self._pick_colors(payload, len(groups))
                    for i, g in enumerate(groups):
                        sub = chart_df[chart_df[color_col].astype(str) == g]
                        sub_sizes = sizes[sub.index] if sizes is not None else None
                        ax.scatter(
                            sub[x_col].astype(float),
                            sub[y_col].astype(float),
                            s=sub_sizes,
                            alpha=0.75,
                            label=str(g),
                            color=colors[i]
                        )
            else:
                color = self._pick_colors(payload, 1)[0]
                ax.scatter(chart_df[x_col].astype(float), chart_df[y_col].astype(float), s=sizes, alpha=0.75, color=color)

            if trendline and len(chart_df) >= 2:
                x_num = chart_df[x_col].astype(float).values
                y_num = chart_df[y_col].astype(float).values
                trend = self._build_linear_trend(x_num, y_num)
                if trend is not None:
                    order = np.argsort(x_num)
                    ax.plot(x_num[order], trend[order], linestyle="--", linewidth=2, label="Tendência")

            ax.set_xlabel(x_label or "")
            ax.set_ylabel(y_label or "")
            ax.xaxis.set_major_formatter(mticker.FuncFormatter(fmt_axis_value_factory(value_format)))
            ax.yaxis.set_major_formatter(mticker.FuncFormatter(fmt_axis_value_factory(value_format)))

        # ---------------------------------------------------------------------
        # INDICATOR
        # ---------------------------------------------------------------------
        elif chart_type == "indicator":
            value_column = value_col or y_col
            agg = str(payload.get("aggregation") or "sum").strip().lower()
            values = chart_df[value_column].astype(float)

            if agg == "avg":
                final_value = float(values.mean())
            elif agg == "min":
                final_value = float(values.min())
            elif agg == "max":
                final_value = float(values.max())
            elif agg == "count":
                final_value = float(values.count())
            else:
                final_value = float(values.sum())

            secondary_value = None
            secondary_value_col = payload.get("secondary_value_col")
            if secondary_value_col and secondary_value_col in chart_df.columns:
                values2 = chart_df[secondary_value_col].astype(float)
                secondary_value = float(values2.sum())

            ax.axis("off")
            color = self._pick_colors(payload, 1)[0]
            suffix = str(payload.get("suffix") or "").strip()
            prefix = str(payload.get("prefix") or "").strip()

            if value_format == "percent":
                main_text = f"{prefix}{self._fmt_ptbr(final_value)}%{suffix}"
            elif value_format == "integer":
                main_text = f"{prefix}{self._fmt_ptbr(final_value, casas=0)}{suffix}"
            elif value_format == "million":
                main_text = f"{prefix}{self._fmt_ptbr(final_value / 1_000_000)} MM{suffix}"
            else:
                main_text = f"{prefix}{self._fmt_ptbr(final_value)}{suffix}"

            ax.text(0.5, 0.60, main_text, ha="center", va="center", fontsize=34, fontweight="bold", color=color)
            ax.text(0.5, 0.42, str(payload.get("indicator_label") or title), ha="center", va="center", fontsize=16)

            if secondary_value is not None:
                if secondary_value_format == "percent":
                    sec_text = f"{self._fmt_ptbr(secondary_value)}%"
                elif secondary_value_format == "integer":
                    sec_text = self._fmt_ptbr(secondary_value, casas=0)
                elif secondary_value_format == "million":
                    sec_text = f"{self._fmt_ptbr(secondary_value / 1_000_000)} MM"
                else:
                    sec_text = self._fmt_ptbr(secondary_value)

                ax.text(0.5, 0.26, f"Valor secundário: {sec_text}", ha="center", va="center", fontsize=11)

        # ---------------------------------------------------------------------
        # CANDLESTICK
        # ---------------------------------------------------------------------
        elif chart_type == "candlestick":
            xdates = pd.to_datetime(chart_df[x_col])
            candle_width = float(payload.get("candle_width") or 0.6)

            up_color = str(payload.get("up_color") or "#2ca02c")
            down_color = str(payload.get("down_color") or "#d62728")

            for _, row in chart_df.iterrows():
                x = mdates.date2num(pd.to_datetime(row[x_col]))
                o = float(row[open_col])
                h = float(row[high_col])
                l = float(row[low_col])
                c = float(row[close_col])

                color = up_color if c >= o else down_color
                ax.vlines(x, l, h, color=color, linewidth=1.2)
                body_bottom = min(o, c)
                body_height = abs(c - o)
                if body_height == 0:
                    body_height = 0.000001
                rect = Rectangle(
                    (x - candle_width / 2, body_bottom),
                    candle_width,
                    body_height,
                    facecolor=color,
                    edgecolor=color,
                    alpha=0.85
                )
                ax.add_patch(rect)

            if trendline and len(chart_df) >= 2:
                trend = self._build_linear_trend(np.arange(len(chart_df), dtype=float), chart_df[close_col].astype(float).values)
                if trend is not None:
                    ax.plot(xdates, trend, linestyle="--", linewidth=2, label="Tendência")

            ax.set_xlabel(x_label or "")
            ax.set_ylabel(y_label or close_col or "")
            ax.yaxis.set_major_formatter(mticker.FuncFormatter(fmt_axis_value_factory(value_format)))
            ax.xaxis_date()
            ax.xaxis.set_major_formatter(mdates.DateFormatter("%Y-%m-%d"))
            fig.autofmt_xdate()

        # ---------------------------------------------------------------------
        # COMBO
        # ---------------------------------------------------------------------
        elif chart_type == "combo":
            work = chart_df.sort_values(by=x_col)
            if np.issubdtype(work[x_col].dtype, np.datetime64):
                x_vals = work[x_col]
                x_plot = work[x_col]
            else:
                x_vals = work[x_col].astype(str)
                x_plot = np.arange(len(work))

            combo_mode = str(payload.get("combo_mode") or "column_line").strip().lower()
            bar_color = self._pick_colors(payload, 2)[0]
            line_color = self._pick_colors(payload, 2)[1]

            if combo_mode in ("bar_line",):
                ax.bar(x_plot, work[y_col].astype(float), color=bar_color, alpha=0.8, label=y_col)
            else:
                ax.bar(x_plot, work[y_col].astype(float), color=bar_color, alpha=0.8, label=y_col)

            ax2 = ax.twinx()
            ax2.plot(x_plot, work[y2_col].astype(float), color=line_color, marker="o", linewidth=2.0, label=y2_col)

            if trendline and len(work) >= 2:
                trend = self._build_linear_trend(np.arange(len(work), dtype=float), work[y2_col].astype(float).values)
                if trend is not None:
                    ax2.plot(x_plot, trend, linestyle="--", linewidth=2, label=f"Tendência {y2_col}")

            if np.issubdtype(work[x_col].dtype, np.datetime64):
                ax.set_xticks(mdates.date2num(pd.to_datetime(work[x_col])))
                ax.set_xticklabels(pd.to_datetime(work[x_col]).dt.strftime("%Y-%m-%d"), rotation=45, ha="right")
            else:
                ax.set_xticks(x_plot)
                ax.set_xticklabels(work[x_col].astype(str), rotation=45, ha="right")

            ax.set_xlabel(x_label or "")
            ax.set_ylabel(y_label or "")
            ax2.set_ylabel(y2_label or "")
            ax.yaxis.set_major_formatter(mticker.FuncFormatter(fmt_axis_value_factory(value_format)))
            ax2.yaxis.set_major_formatter(mticker.FuncFormatter(fmt_axis_value_factory(secondary_value_format)))

            if show_legend:
                handles1, labels1 = ax.get_legend_handles_labels()
                handles2, labels2 = ax2.get_legend_handles_labels()
                ax.legend(handles1 + handles2, labels1 + labels2, loc="best")

        # ---------------------------------------------------------------------
        # labels / legend / title
        # ---------------------------------------------------------------------
        if show_labels and chart_type in ("bar", "column") and not series_col:
            for container in ax.containers:
                try:
                    ax.bar_label(container, fmt=lambda v: fmt_axis_value_factory(value_format)(v))
                except Exception:
                    pass

        if title:
            ax.set_title(title, fontsize=16, fontweight="bold", pad=16)

        if subtitle:
            fig.text(0.5, 0.94, subtitle, ha="center", fontsize=10)

        if show_legend and chart_type not in ("pie", "indicator", "combo"):
            handles, labels = ax.get_legend_handles_labels()
            if handles and labels:
                ax.legend()

        if chart_type not in ("indicator",):
            ax.grid(axis="y", alpha=0.25)

        fig.tight_layout()

        ts = time.strftime("%Y%m%d_%H%M%S")
        fname = f"{self._slugify(title)}_{ts}.png"
        local_path = os.path.join(os.getcwd(), fname)
        fig.savefig(local_path, dpi=dpi, bbox_inches="tight")
        plt.close(fig)

        return local_path

    # =========================================================================
    # Excel
    # =========================================================================

    def _write_excel(self, df: pd.DataFrame, label: str) -> str:
        ts = time.strftime("%Y%m%d_%H%M%S")
        safe_file_name = self._slugify(label)
        file_name = f"{safe_file_name}_{ts}.xlsx"
        local_path = os.path.join(os.getcwd(), file_name)

        sheet_name = self._sheet_name(label)

        df_to_save = df.copy()
        for col in df_to_save.columns:
            if pd.api.types.is_datetime64_any_dtype(df_to_save[col]):
                df_to_save[col] = pd.to_datetime(df_to_save[col], errors="coerce").dt.strftime("%Y-%m-%d %H:%M:%S")

        with pd.ExcelWriter(local_path, engine="openpyxl") as writer:
            df_to_save.to_excel(writer, index=False, sheet_name=sheet_name)

        wb = load_workbook(local_path)
        ws = wb[sheet_name]

        bold_font = Font(bold=True)
        for cell in ws[1]:
            cell.font = bold_font

        for col_cells in ws.columns:
            max_length = 0
            col_letter = col_cells[0].column_letter
            for cell in col_cells:
                try:
                    val = "" if cell.value is None else str(cell.value)
                    if len(val) > max_length:
                        max_length = len(val)
                except Exception:
                    pass
            ws.column_dimensions[col_letter].width = min(max(max_length + 2, 12), 60)

        wb.save(local_path)
        wb.close()

        return local_path

    # =========================================================================
    # S3
    # =========================================================================

    def _upload_graph_s3_and_presign(self, local_path: str) -> Tuple[str, str]:
        bucket, base_prefix = self._split_s3_uri(self.graphs_s3_base)
        file_name = os.path.basename(local_path)
        key = base_prefix.rstrip("/") + "/" + file_name if base_prefix else file_name

        self.s3_client.upload_file(local_path, bucket, key)

        try:
            os.remove(local_path)
        except Exception:
            pass

        graph_s3_url = f"s3://{bucket}/{key}"

        graph_http_url = self.s3_client.generate_presigned_url(
            ClientMethod="get_object",
            Params={"Bucket": bucket, "Key": key},
            ExpiresIn=self.presign_expires_seconds,
        )

        return graph_s3_url, graph_http_url

    def _upload_excel_s3_and_presign(self, local_path: str) -> Tuple[str, str]:
        bucket, base_prefix = self._split_s3_uri(self.excel_s3_base)
        file_name = os.path.basename(local_path)
        key = base_prefix.rstrip("/") + "/" + file_name if base_prefix else file_name

        self.s3_client.upload_file(local_path, bucket, key)

        try:
            os.remove(local_path)
        except Exception:
            pass

        excel_s3_url = f"s3://{bucket}/{key}"

        excel_http_url = self.s3_client.generate_presigned_url(
            ClientMethod="get_object",
            Params={"Bucket": bucket, "Key": key},
            ExpiresIn=self.presign_expires_seconds,
        )

        return excel_s3_url, excel_http_url

    # =========================================================================
    # Texto
    # =========================================================================

    def _build_summary_txt(
        self,
        chart_type: str,
        row_count_raw: int,
        row_count_chart: int,
        columns_raw: List[str],
        columns_chart: List[str],
        title: str,
    ) -> str:
        return (
            f"Gráfico '{title}' gerado com sucesso. "
            f"Tipo: {chart_type}. "
            f"Linhas retornadas pela SQL: {row_count_raw}. "
            f"Linhas efetivamente usadas no gráfico: {row_count_chart}. "
            f"Colunas retornadas pela SQL: {', '.join(columns_raw)}. "
            f"Colunas exportadas no Excel: {', '.join(columns_chart)}."
        )

    # =========================================================================
    # Tool por payload JSON
    # =========================================================================

    def _run_chart_payload(self) -> Callable[[Union[str, Dict[str, Any]]], str]:
        """
        Entrada:
        {
          "sql": "...",
          "chart_type": "pie|bar|column|line|area|scatter|indicator|candlestick|combo",
          "title": "Meu gráfico",
          "subtitle": "Opcional",

          "x_col": "mes",
          "y_col": "valor",
          "y2_col": "meta",

          "series_col": "categoria",
          "category_col": "categoria",
          "value_col": "valor",
          "size_col": "tam",
          "color_col": "grupo",

          "open_col": "open",
          "high_col": "high",
          "low_col": "low",
          "close_col": "close",

          "trendline": true,
          "stacked": false,
          "horizontal": false,
          "show_legend": true,
          "show_labels": false,
          "donut": false,

          "top_n": 10,
          "sort_by": "valor",
          "sort_ascending": false,

          "colors": ["#1f77b4", "#ff7f0e"],

          "x_label": "Mês",
          "y_label": "Valor",
          "y2_label": "Meta",

          "value_format": "number|currency|percent|integer|million",
          "secondary_value_format": "number|currency|percent|integer|million",

          "figsize": [14, 8],
          "dpi": 170,

          "aggregation": "sum|avg|min|max|count",
          "prefix": "R$ ",
          "suffix": "",
          "indicator_label": "Lucro líquido",

          "combo_mode": "column_line|bar_line",
          "secondary_value_col": "valor2",

          "candle_width": 0.6,
          "up_color": "#2ca02c",
          "down_color": "#d62728"
        }
        """
        def _tool(input_payload: Union[str, Dict[str, Any]]) -> str:
            # 0) parse
            if input_payload is None:
                return json.dumps(
                    [{"_meta": {"ok": False, "error": "Nenhuma entrada foi fornecida (esperado JSON)."}}],
                    ensure_ascii=False,
                    indent=2,
                )

            if isinstance(input_payload, str):
                if not input_payload.strip():
                    return json.dumps(
                        [{"_meta": {"ok": False, "error": "Entrada vazia (esperado JSON)."}}],
                        ensure_ascii=False,
                        indent=2,
                    )
                try:
                    payload = json.loads(input_payload)
                except Exception as e:
                    return json.dumps(
                        [{"_meta": {"ok": False, "error": f"Entrada não é JSON válido: {e}"}}],
                        ensure_ascii=False,
                        indent=2,
                    )
            elif isinstance(input_payload, dict):
                payload = input_payload
            else:
                return json.dumps(
                    [{"_meta": {"ok": False, "error": "Entrada deve ser JSON string ou dict."}}],
                    ensure_ascii=False,
                    indent=2,
                )

            # 1) validações básicas
            sql = str(payload.get("sql") or "").strip()
            chart_type = self._normalize_chart_type(payload.get("chart_type"))
            title = str(payload.get("title") or payload.get("label") or f"grafico_{chart_type}").strip()

            if not sql:
                return json.dumps(
                    [{"_meta": {"ok": False, "error": "SQL vazia ou inválida."}}],
                    ensure_ascii=False,
                    indent=2,
                )

            if chart_type not in self.SUPPORTED_CHART_TYPES:
                return json.dumps(
                    [{"_meta": {
                        "ok": False,
                        "error": f"chart_type inválido: {chart_type}. Use: {sorted(self.SUPPORTED_CHART_TYPES)}"
                    }}],
                    ensure_ascii=False,
                    indent=2,
                )

            # 2) query
            try:
                raw_df = self._execute_athena_query(sql)
            except Exception as e:
                return json.dumps(
                    [{"_meta": {"ok": False, "stage": "athena_query", "error": str(e)}}],
                    ensure_ascii=False,
                    indent=2,
                )

            if raw_df.empty:
                return json.dumps(
                    [{"_meta": {"ok": False, "stage": "athena_query", "error": "A SQL não retornou linhas."}}],
                    ensure_ascii=False,
                    indent=2,
                )

            # 3) preparar dados usados no gráfico
            try:
                chart_df = self._prepare_chart_dataframe(raw_df, payload)
            except Exception as e:
                return json.dumps(
                    [{"_meta": {"ok": False, "stage": "prepare_chart_data", "error": str(e)}}],
                    ensure_ascii=False,
                    indent=2,
                )

            if chart_df.empty:
                return json.dumps(
                    [{"_meta": {"ok": False, "stage": "prepare_chart_data", "error": "Não sobraram dados válidos para gerar o gráfico."}}],
                    ensure_ascii=False,
                    indent=2,
                )

            # 4) gráfico local
            try:
                graph_local_path = self._create_chart_image(chart_df=chart_df, payload=payload)
            except Exception as e:
                return json.dumps(
                    [{"_meta": {"ok": False, "stage": "create_chart", "error": str(e)}}],
                    ensure_ascii=False,
                    indent=2,
                )

            # 5) excel local
            try:
                excel_local_path = self._write_excel(df=chart_df, label=title)
            except Exception as e:
                return json.dumps(
                    [{"_meta": {"ok": False, "stage": "write_excel", "error": str(e)}}],
                    ensure_ascii=False,
                    indent=2,
                )

            # 6) upload gráfico
            try:
                graph_s3_url, graph_http_url = self._upload_graph_s3_and_presign(graph_local_path)
            except Exception as e:
                return json.dumps(
                    [{"_meta": {"ok": False, "stage": "upload_graph_s3", "error": str(e)}}],
                    ensure_ascii=False,
                    indent=2,
                )

            # 7) upload excel
            try:
                excel_s3_url, excel_http_url = self._upload_excel_s3_and_presign(excel_local_path)
            except Exception as e:
                return json.dumps(
                    [{"_meta": {"ok": False, "stage": "upload_excel_s3", "error": str(e)}}],
                    ensure_ascii=False,
                    indent=2,
                )

            # 8) summary
            summary_txt = self._build_summary_txt(
                chart_type=chart_type,
                row_count_raw=int(raw_df.shape[0]),
                row_count_chart=int(chart_df.shape[0]),
                columns_raw=list(raw_df.columns),
                columns_chart=list(chart_df.columns),
                title=title,
            )

            meta = {
                "ok": True,
                "chart_type": chart_type,
                "title": title,
                "row_count_raw": int(raw_df.shape[0]),
                "col_count_raw": int(raw_df.shape[1]),
                "columns_raw": list(raw_df.columns),
                "row_count_chart": int(chart_df.shape[0]),
                "col_count_chart": int(chart_df.shape[1]),
                "columns_chart": list(chart_df.columns),
                "graph_s3_url": graph_s3_url,
                "graph_http_url": graph_http_url,
                "excel_s3_url": excel_s3_url,
                "excel_http_url": excel_http_url,
                "summary_txt": summary_txt,
            }

            return json.dumps(
                [{"_meta": meta}],
                ensure_ascii=False,
                indent=2,
                default=str,
            )

        return _tool

    # =========================================================================
    # Tool pública
    # =========================================================================

    def run_chart_tool(self) -> Callable[[Union[str, Dict[str, Any]]], str]:
        """
        Tool pública no formato payload JSON/dict.
        """
        return self._run_chart_payload()
    
